"""Approval workflow, task tracking, overdue maths, dashboard, CSV and report downloads."""
from datetime import timedelta

from django.utils import timezone

from meetings.models import Decision, Question
from tasks.models import Task

from .base import BaseAPITest


class ReviewWorkflowTests(BaseAPITest):
    def setUp(self):
        super().setUp()
        self.meeting = self.make_meeting()
        self.analyze(self.meeting)
        self.task = Task.objects.get(title="Send revised proposal")

    def test_only_approved_tasks_in_operational_list(self):
        self.assertEqual(self.client.get("/api/tasks/").data["count"], 0)
        self.assertEqual(self.client.get("/api/tasks/?view=review").data["count"], 3)
        self.client.post(f"/api/tasks/{self.task.pk}/approve/")
        operational = self.client.get("/api/tasks/").data
        self.assertEqual([t["title"] for t in operational["results"]], ["Send revised proposal"])
        self.assertEqual(self.client.get("/api/tasks/?view=review").data["count"], 2)

    def test_reject_draft_task(self):
        r = self.client.post(f"/api/tasks/{self.task.pk}/reject/")
        self.assertEqual(r.data["review_status"], "rejected")
        self.assertEqual(self.client.get("/api/tasks/").data["count"], 0)
        self.assertEqual(self.client.get("/api/tasks/?view=rejected").data["count"], 1)
        self.assertEqual(self.client.post(f"/api/tasks/{self.task.pk}/reset/").data["review_status"], "needs_confirmation")

    def test_approve_and_reject_decision_risk_question(self):
        d = Decision.objects.get()
        self.assertEqual(self.client.post(f"/api/decisions/{d.pk}/approve/").data["review_status"], "approved")
        q = Question.objects.get()
        self.assertEqual(self.client.post(f"/api/questions/{q.pk}/reject/").data["review_status"], "rejected")
        from meetings.models import Risk
        rk = Risk.objects.filter(risk_type="explicit").get()
        self.assertEqual(self.client.post(f"/api/risks/{rk.pk}/approve/").data["review_status"], "approved")

    def test_editing_marks_user_edited_without_approving(self):
        d = Decision.objects.get()
        r = self.client.patch(f"/api/decisions/{d.pk}/", {"text": "Launch on staging, then production"}, format="json")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data["user_edited"])
        self.assertNotEqual(r.data["review_status"], "approved")
        # review_status and evidence cannot be changed through PATCH
        r = self.client.patch(f"/api/decisions/{d.pk}/", {"review_status": "approved", "evidence": "x"}, format="json")
        self.assertNotEqual(r.data["review_status"], "approved")
        self.assertNotEqual(r.data["evidence"], "x")

    def test_bulk_approval_requires_confirmation(self):
        items = [{"type": "task", "id": t.pk} for t in Task.objects.all()]
        url = f"/api/meetings/{self.meeting.pk}/review/"
        r = self.client.post(url, {"action": "approve", "items": items}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.data["error"]["code"], "confirmation_required")
        self.assertFalse(Task.objects.filter(review_status="approved").exists())
        r = self.client.post(url, {"action": "approve", "confirm": True, "items": items}, format="json")
        self.assertEqual(r.data["updated"], 3)
        self.assertEqual(Task.objects.filter(review_status="approved").count(), 3)

    def test_nothing_is_auto_approved(self):
        self.assertFalse(Task.objects.filter(review_status="approved").exists())
        self.assertFalse(Decision.objects.filter(review_status="approved").exists())

    def test_cannot_track_status_before_approval(self):
        r = self.client.patch(f"/api/tasks/{self.task.pk}/", {"status": "in_progress"}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("status", r.data["error"]["details"])

    def test_followup_rebuild_uses_only_approved_records(self):
        self.client.post(f"/api/tasks/{self.task.pk}/approve/")
        r = self.client.post(f"/api/meetings/{self.meeting.pk}/followup/rebuild/")
        text = r.data["analysis"]["followup_draft"]
        self.assertIn("Send revised proposal", text)
        self.assertNotIn("Migrate the database", text)
        edited = self.client.patch(f"/api/meetings/{self.meeting.pk}/analysis/", {"followup_draft": "My edit"}, format="json")
        self.assertEqual(edited.data["analysis"]["followup_draft"], "My edit")


class TaskTrackingTests(BaseAPITest):
    def mk(self, title, **kw):
        defaults = dict(user=self.user, meeting=None, review_status="approved")
        defaults.update(kw)
        return Task.objects.create(title=title, **defaults)

    def test_overdue_calculation(self):
        today = timezone.localdate()
        past = self.mk("past", due_date=today - timedelta(days=2))
        future = self.mk("future", due_date=today + timedelta(days=2))
        no_date = self.mk("no date", due_date=None)
        done = self.mk("done", due_date=today - timedelta(days=5), status="completed")
        draft = self.mk("draft", due_date=today - timedelta(days=5), review_status="draft")
        self.assertTrue(past.is_overdue)
        for task in (future, no_date, done, draft):
            self.assertFalse(task.is_overdue, task.title)  # absent deadline is never overdue
        overdue = self.client.get("/api/tasks/?overdue=true").data["results"]
        self.assertEqual([t["title"] for t in overdue], ["past"])
        self.assertEqual(self.client.get("/api/dashboard/stats/").data["overdue_count"], 1)

    def test_status_updates_and_completion_timestamp(self):
        task = self.mk("t")
        for status in ("in_progress", "blocked", "completed"):
            r = self.client.patch(f"/api/tasks/{task.pk}/", {"status": status}, format="json")
            self.assertEqual(r.data["status"], status)
        task.refresh_from_db()
        self.assertIsNotNone(task.completed_at)
        self.client.patch(f"/api/tasks/{task.pk}/", {"status": "pending"}, format="json")
        task.refresh_from_db()
        self.assertIsNone(task.completed_at)
        self.assertEqual(self.client.patch(f"/api/tasks/{task.pk}/", {"status": "nonsense"}, format="json").status_code, 400)

    def test_manual_task_create_edit_delete(self):
        r = self.client.post("/api/tasks/", {"title": "Call vendor", "priority": "high"}, format="json")
        self.assertEqual(r.status_code, 201)
        self.assertEqual((r.data["review_status"], r.data["owner_name"], r.data["due_date"]), ("approved", None, None))
        self.assertEqual(self.client.post("/api/tasks/", {"title": "  "}, format="json").status_code, 400)
        self.assertEqual(self.client.patch(f"/api/tasks/{r.data['id']}/", {"owner_name": "  "}, format="json").data["owner_name"], None)
        self.assertEqual(self.client.delete(f"/api/tasks/{r.data['id']}/").status_code, 204)

    def test_filters_search_sorting_and_pagination(self):
        today = timezone.localdate()
        self.mk("Write spec", owner_name="Priya", priority="high", due_date=today + timedelta(days=3))
        self.mk("Review budget", owner_name="Rahul", priority="low", due_date=today + timedelta(days=1), status="blocked")
        self.mk("Plan launch", owner_name=None, priority=None, due_date=None)
        get = lambda q: [t["title"] for t in self.client.get(f"/api/tasks/?{q}").data["results"]]
        self.assertEqual(get("search=budget"), ["Review budget"])
        self.assertEqual(get("status=blocked"), ["Review budget"])
        self.assertEqual(get("priority=high"), ["Write spec"])
        self.assertEqual(get("owner=priya"), ["Write spec"])
        self.assertEqual(get("owner=unassigned"), ["Plan launch"])
        self.assertEqual(get("no_due_date=true"), ["Plan launch"])
        self.assertEqual(get(f"due_to={today + timedelta(days=2)}"), ["Review budget"])
        self.assertEqual(get("ordering=due_date"), ["Review budget", "Write spec", "Plan launch"])  # no date last
        self.assertEqual(get("ordering=-priority")[0], "Write spec")
        self.assertEqual(self.client.get("/api/tasks/?page_size=2").data["total_pages"], 2)

    def test_csv_export(self):
        self.mk("Normal task", owner_name="Priya")
        self.mk("=HYPERLINK(\"http://evil\")", owner_name=None)
        r = self.client.get("/api/tasks/export/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("text/csv", r["Content-Type"])
        self.assertIn("attachment", r["Content-Disposition"])
        text = r.content.decode("utf-8-sig")
        self.assertIn("Normal task", text)
        self.assertIn("Unassigned", text)
        self.assertIn("Not specified", text)
        self.assertIn("'=HYPERLINK", text)  # formula injection neutralised
        chosen = self.client.get(f"/api/tasks/export/?ids={Task.objects.get(title='Normal task').pk}")
        self.assertNotIn("HYPERLINK", chosen.content.decode("utf-8-sig"))


class DashboardAndReportTests(BaseAPITest):
    def test_new_account_has_zero_metrics_and_no_sample_data(self):
        stats = self.client.get("/api/dashboard/stats/").data
        self.assertEqual(stats["total_meetings"], 0)
        self.assertEqual((stats["open_task_count"], stats["overdue_count"], stats["meetings_awaiting_review"]), (0, 0, 0))
        self.assertEqual(stats["recent_meetings"], [])
        self.assertEqual(stats["upcoming_deadlines"], [])

    def test_metrics_come_from_real_records(self):
        meeting = self.make_meeting()
        self.analyze(meeting)
        task = Task.objects.get(title="Send revised proposal")
        self.client.post(f"/api/tasks/{task.pk}/approve/")
        stats = self.client.get("/api/dashboard/stats/").data
        self.assertEqual(stats["total_meetings"], 1)
        self.assertEqual(stats["meetings_awaiting_review"], 1)
        self.assertEqual(stats["open_task_count"], 1)
        self.assertEqual(stats["task_status_counts"]["pending"], 1)
        self.client.patch(f"/api/tasks/{task.pk}/", {"status": "completed"}, format="json")
        self.assertEqual(self.client.get("/api/dashboard/stats/").data["open_task_count"], 0)

    def test_meeting_report_downloads(self):
        meeting = self.make_meeting()
        self.analyze(meeting)
        pdf = self.client.get(f"/api/meetings/{meeting.pk}/report/?file=pdf")
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        docx = self.client.get(f"/api/meetings/{meeting.pk}/report/?file=docx")
        self.assertTrue(docx.content.startswith(b"PK"))
        self.assertEqual(self.client.get(f"/api/meetings/{meeting.pk}/report/?file=exe").status_code, 400)

    def test_report_excludes_unapproved_drafts(self):
        import io
        from docx import Document

        meeting = self.make_meeting()
        self.analyze(meeting)
        task = Task.objects.get(title="Send revised proposal")
        self.client.post(f"/api/tasks/{task.pk}/approve/")
        doc = Document(io.BytesIO(self.client.get(f"/api/meetings/{meeting.pk}/report/?file=docx").content))
        text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn("Send revised proposal", text)
        self.assertNotIn("Migrate the database", text)

    def test_questions_export_and_health_endpoint(self):
        meeting = self.make_meeting()
        self.analyze(meeting)
        q = Question.objects.get()
        self.client.post(f"/api/questions/{q.pk}/approve/")
        text = self.client.get("/api/reports/questions/export/").content.decode("utf-8-sig")
        self.assertIn("Has the client approved the budget?", text)
        from rest_framework.test import APIClient
        self.assertEqual(APIClient().get("/api/health/").data, {"status": "ok"})

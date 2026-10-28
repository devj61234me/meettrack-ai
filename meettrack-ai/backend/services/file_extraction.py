"""
Extract plain text from uploaded TXT / DOCX / text-based PDF files.
Files are processed in memory and never written to disk.
"""
import io
import os
import re
import zipfile

from django.conf import settings

ALLOWED_EXTENSIONS = (".txt", ".docx", ".pdf")
MAX_PDF_PAGES = 200
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024  # protects against zip bombs


class FileExtractionError(Exception):
    """Safe-to-display message about why a file could not be read."""


def _clean(text):
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _read_txt(data):
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileExtractionError("The text file uses an unsupported encoding. Save it as UTF-8 and try again.")


def _read_docx(data):
    if not data.startswith(b"PK"):
        raise FileExtractionError("This file is not a valid .docx document.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise FileExtractionError("This document is too large to process safely.")
        from docx import Document

        document = Document(io.BytesIO(data))
    except FileExtractionError:
        raise
    except Exception:  # noqa: BLE001 - corrupt or unsupported file
        raise FileExtractionError("This .docx file could not be read. It may be corrupted or password protected.")
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _read_pdf(data):
    if not data.startswith(b"%PDF"):
        raise FileExtractionError("This file is not a valid PDF document.")
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise FileExtractionError("This PDF is password protected. Remove the password and try again.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise FileExtractionError(f"This PDF has more than {MAX_PDF_PAGES} pages. Please upload a shorter file.")
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except FileExtractionError:
        raise
    except Exception:  # noqa: BLE001
        raise FileExtractionError("This PDF could not be read. It may be corrupted.")
    return text


def extract_text(uploaded_file):
    """Return (clean_text, safe_filename). Raises FileExtractionError with a user-friendly message."""
    name = os.path.basename(uploaded_file.name or "")
    extension = os.path.splitext(name)[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise FileExtractionError("Unsupported file type. Please upload a .txt, .docx or text-based .pdf file.")
    max_bytes = settings.MAX_UPLOAD_MB * 1024 * 1024
    if uploaded_file.size > max_bytes:
        raise FileExtractionError(f"The file is larger than the {settings.MAX_UPLOAD_MB} MB limit.")
    data = uploaded_file.read()
    if not data:
        raise FileExtractionError("The file is empty.")

    reader = {".txt": _read_txt, ".docx": _read_docx, ".pdf": _read_pdf}[extension]
    text = _clean(reader(data))
    if len(text) < 20:
        if extension == ".pdf":
            raise FileExtractionError(
                "No readable text was found in this PDF. Scanned or image-only PDFs are not supported."
            )
        raise FileExtractionError("No readable text was found in this file.")
    if len(text) > settings.MAX_NOTES_CHARS:
        raise FileExtractionError(
            f"The extracted text is longer than {settings.MAX_NOTES_CHARS:,} characters. Please split the notes."
        )
    return text, name[:255]
